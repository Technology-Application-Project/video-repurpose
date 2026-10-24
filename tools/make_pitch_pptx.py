"""Builds docs/pitch/brevio-pitch.pptx from the pitch content, with the narration as speaker notes.

Usage: python tools/make_pitch_pptx.py   (needs python-pptx; run tools/make_pitch_assets.py first for timings)
"""
from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
PITCH = ROOT / "docs" / "pitch"
STILLS = PITCH / "stills"

SERIF, SANS = "Georgia", "Calibri"
C = {"green": "145f4c", "gold": "d7ac40", "ink": "202d28", "muted": "556760", "line": "d5dfda", "red": "a34e48",
     "paper": "f7f9f8", "mint": "eef5f1", "mintline": "b8cec3", "goldtint": "fcf5df", "goldline": "d8bd69",
     "redtint": "f6e6e4", "redline": "d9a8a3", "slate": "2d4a5a", "white": "ffffff", "darkgold": "8a6a10", "track": "e6eeea"}


def rgb(name: str) -> RGBColor:
    return RGBColor.from_string(C.get(name, name))


def box(slide, x, y, w, h, fill=None, line=None, lw=1.25, rounded=True, radius=0.08):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.shadow.inherit = False
    if rounded:
        shape.adjustments[0] = radius
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(lw)
    else:
        shape.line.fill.background()
    return shape


def _fill_frame(tf, content, size, bold, color, font, align, italic):
    paragraphs = content if isinstance(content, list) else content.split("\n")
    for index, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.alignment = align
        runs = para if isinstance(para, list) else [(para, {})]
        for text, style in runs:
            r = p.add_run()
            r.text = text
            r.font.size = Pt(style.get("size", size))
            r.font.bold = style.get("bold", bold)
            r.font.italic = italic
            r.font.name = style.get("font", font)
            r.font.color.rgb = rgb(style.get("color", color))


def txt(slide, x, y, w, h, content, size=16, bold=False, color="ink", font=SANS, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    _fill_frame(tf, content, size, bold, color, font, align, italic)
    return tb


def pill(slide, x, y, w, h, text, fill="mint", line="mintline", color="green", size=14, bold=True):
    shape = box(slide, x, y, w, h, fill, line, radius=0.5)
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = True
    _fill_frame(tf, text, size, bold, color, SANS, PP_ALIGN.CENTER, False)
    return shape


def card(slide, x, y, w, h, fill="white", line="line", lw=1.5):
    return box(slide, x, y, w, h, fill, line, lw)


def picture(slide, name, x, y, w):
    pic = slide.shapes.add_picture(str(STILLS / name), Inches(x), Inches(y), width=Inches(w))
    pic.line.color.rgb = rgb("mintline")
    pic.line.width = Pt(1.5)
    return pic


def mmss(seconds: float) -> str:
    return f"{int(seconds // 60)}:{int(seconds % 60):02d}"


def load_timings() -> dict:
    raw = (PITCH / "timings.js").read_text(encoding="utf-8")
    return {scene["id"]: scene for scene in json.loads(raw[raw.index("{"): raw.rindex("}") + 1])["scenes"]}


class Deck:
    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.timings = load_timings()
        self.count = 0

    def slide(self, label: str, title: str | None, cue: tuple[str, int, int] | None, source: str = "", notes: str = ""):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.count += 1
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb("paper")
        box(slide, 0, 0, 13.333, 0.09, "green", rounded=False)
        txt(slide, 0.6, 0.28, 8, 0.3, f"{self.count:02d} \u00b7 {label.upper()}", 12, True, "green")
        timing = ""
        narration = ""
        if cue:
            scene = self.timings[cue[0]]
            first, last = scene["sentences"][cue[1]], scene["sentences"][cue[2]]
            begin = scene["start"] if cue[1] == 0 else first["start"]
            timing = f"{mmss(begin)}\u2013{mmss(last['end'])}"
            narration = " ".join(s["text"] for s in scene["sentences"][cue[1]: cue[2] + 1])
        if timing:
            txt(slide, 10.2, 0.28, 2.53, 0.3, timing, 12, True, "muted", align=PP_ALIGN.RIGHT)
        if title:
            txt(slide, 0.6, 0.72, 12.1, 0.9, title, 28, True, "ink", SERIF, anchor=MSO_ANCHOR.MIDDLE)
        if source:
            txt(slide, 0.6, 7.0, 12.1, 0.38, source, 9.5, False, "muted")
        body = f"[{timing}] Say:\n{narration}" if narration else ""
        slide.notes_slide.notes_text_frame.text = (body + ("\n\n" + notes if notes else "")).strip()
        return slide


def build() -> Path:
    d = Deck()

    # 1 Hook
    s = d.slide("Hook", None, ("hook", 0, 2), "", "On screen: a 30:00 lecture bar condenses into a roughly 40 s short video; the title lands on the last sentence.")
    txt(s, 0.9, 0.55, 8, 0.4, [[("30:00", {"bold": True, "size": 20, "color": "ink"}), ("   one lecture, 1,800 seconds", {})]], 16, False, "muted")
    selected = {9, 10, 11, 12, 30, 31, 32, 33, 34, 35, 52, 53, 54}
    for k in range(64):
        h = 0.35 + ((k * 37 + (k % 5) * 19) % 70) / 100 * 0.95
        box(s, 0.9 + k * 0.18, 1.05 + (1.4 - h) / 2, 0.13, h, "gold" if k in selected else "mintline", rounded=False)
    txt(s, 0.9, 2.75, 8, 0.4, [[("about 0:40", {"bold": True, "size": 20, "color": "ink"}), ("   one short video for the audience you choose", {})]], 16, False, "muted")
    parts = [("Opening", 1, False), ("Clip", 2.4, True), ("Bridge", .9, False), ("Clip", 2.4, True), ("Bridge", .9, False), ("Clip", 2.4, True), ("Ending", 1, False)]
    unit, x = (11.5 - 0.08 * 6) / sum(p[1] for p in parts), 0.9
    for name, weight, is_clip in parts:
        pill(s, x, 3.25, weight * unit, 0.62, name, "e3ecf0" if is_clip else "mint", "a8bcc7" if is_clip else "mintline", "slate" if is_clip else "green", 14)
        x += weight * unit + 0.08
    pill(s, 9.1, 4.0, 3.3, 0.45, "Approved by a person", "goldtint", "goldline", "ink", 14)
    txt(s, 0.9, 4.75, 11.5, 1.4, "Brevio Studio", 66, True, "green", SERIF, PP_ALIGN.CENTER)
    txt(s, 0.9, 6.0, 11.5, 0.6, "One long lecture. One short, reviewed video.", 24, False, "muted", align=PP_ALIGN.CENTER)

    # 2 Problem
    s = d.slide("Problem", "Great lectures. Too long to watch. Too risky to automate.", ("problem", 0, 5),
                "Sources: Guo, Kim & Rubin, L@S 2014 (doi 10.1145/2556325.2566239); Wyzowl, State of Video Marketing 2026 (266 respondents); TechSmith, 2026 Video Statistics.",
                "On screen: three stat cards, then the two failure modes, then the conclusion.")
    stats = [("6.9M", "green", "video-watching sessions in an edX study", "Shorter videos are much more engaging", "Guo, Kim & Rubin, L@S 2014"),
             ("71%", "green", "of video marketers say 30 seconds to 2 minutes works best", "", "Wyzowl, State of Video Marketing 2026 (n=266)"),
             ("90%", "red", "of viewers are concerned about AI content, including inaccuracies", "", "TechSmith, 2026 Video Statistics")]
    for i, (num, colour, line1, line2, cite) in enumerate(stats):
        x = 0.6 + i * 4.1
        card(s, x, 1.85, 3.9, 2.85)
        txt(s, x + 0.3, 1.98, 3.4, 1.0, num, 54, True, colour, SERIF)
        txt(s, x + 0.3, 3.05, 3.3, 0.8, line1, 15)
        if line2:
            txt(s, x + 0.3, 3.65, 3.3, 0.5, line2, 15, True)
        txt(s, x + 0.3, 4.2, 3.4, 0.4, cite, 11, False, "muted")
    box(s, 0.6, 4.95, 3.2, 0.75, "goldtint", "goldline").text_frame.text = ""
    txt(s, 0.8, 4.95, 2.9, 0.75, "By hand: slow", 20, True, anchor=MSO_ANCHOR.MIDDLE)
    box(s, 4.0, 4.95, 6.6, 0.75, "redtint", "redline")
    txt(s, 4.2, 4.95, 6.2, 0.75, "Fully automatic: chases virality, hard to trust", 20, True, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, 0.6, 5.95, 12.1, 0.8, [[("Teams are stuck between ", {}), ("slow", {"color": "red"}), (" and ", {}), ("untrustworthy", {"color": "red"}), (".", {})]], 30, True, "ink", SERIF)

    # 3 Process
    s = d.slide("Solution and process", "Four steps, with a human gate before export", ("solution", 0, 1),
                "", "Say the next line of narration while moving to slides 4 and 5, which show the real app.")
    steps = [("1", "Source", "Pick the lecture and the audience: researcher, general public or kids"),
             ("2", "Clips", "Local transcript, strongest clips, and a story plan: opening, bridges, ending"),
             ("3", "Studio", "Narrated preview with captions, music and illustrations, each beside its source evidence"),
             ("4", "Export", "A technical check and a structure check: does the story follow the plan?")]
    for i, (n, name, desc) in enumerate(steps):
        x = 0.6 + i * 3.07
        card(s, x, 1.85, 2.9, 2.25)
        pill(s, x + 0.2, 2.0, 0.5, 0.5, n, "mint", "mintline", "green", 18)
        txt(s, x + 0.85, 2.02, 1.9, 0.5, name, 22, True, "green", SERIF, anchor=MSO_ANCHOR.MIDDLE)
        txt(s, x + 0.2, 2.7, 2.5, 1.3, desc, 14, False, "muted")
    pill(s, 6.2, 1.62, 2.1, 0.36, "Human approval", "goldtint", "goldline", "ink", 12)
    for i, (name, fill, line, colour) in enumerate([("Plan", "mint", "mintline", "green"), ("Build", "mint", "mintline", "green"), ("Check", "mint", "mintline", "green"), ("Approve", "goldtint", "goldline", "ink")]):
        pill(s, 0.6 + i * 1.95, 4.4, 1.6, 0.55, name, fill, line, colour, 18)
        if i < 3:
            txt(s, 2.25 + i * 1.95, 4.4, 0.3, 0.55, "\u2192", 18, False, "muted", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    card(s, 0.6, 5.3, 5.9, 1.4, "mint", "mintline")
    txt(s, 0.85, 5.42, 5.4, 1.2, [[("Technical check", {"bold": True, "size": 17, "color": "green"})], "Media, captions and loudness measured on the exported file."], 14)
    card(s, 6.75, 5.3, 5.98, 1.4, "mint", "mintline")
    txt(s, 7.0, 5.42, 5.5, 1.2, [[("Structure check", {"bold": True, "size": 17, "color": "green"})], "Order, bridges, narration and the plan: does the finished story follow it?"], 14)

    # 4 Demo 1
    s = d.slide("Demo: the real app (1 of 2)", "From source to a planned story", ("solution", 2, 2),
                "Screens: Brevio Studio proof of concept, captured at 2x in a scripted demo session (frontend/scripts/capture-pitch-stills.mjs). Plan data mocked, no cloud calls; lecture picture blurred.",
                "On screen: the real UI. Point at the audience picker, then the blueprint panel.")
    picture(s, "ui-source.png", 0.6, 1.9, 5.95)
    picture(s, "ui-blueprint.png", 6.78, 1.9, 5.95)
    pill(s, 0.6, 5.45, 5.95, 0.5, "1  Pick the lecture and the audience", "goldtint", "goldline", "ink", 16)
    pill(s, 6.78, 5.45, 5.95, 0.5, "2  Plan first: opening, bridges and ending", "goldtint", "goldline", "ink", 16)
    txt(s, 0.6, 6.2, 12.1, 0.6, "The blueprint is written before any narration, so the finished video can be checked against it.", 16, False, "muted", align=PP_ALIGN.CENTER)

    # 5 Demo 2
    s = d.slide("Demo: the real app (2 of 2)", "Review, approve, check", ("solution", 3, 8),
                "Screens: Brevio Studio proof of concept, captured at 2x in a scripted demo session. Export checks are from an existing export; lecture picture blurred.",
                "On screen: Studio preview, the named-reviewer approval, then the export checks.")
    for i, (img, head, desc) in enumerate([("ui-studio.png", "Studio", "Narrated preview, story sequence and source evidence"),
                                           ("ui-approval.png", "Approval", "A named reviewer approves the exact preview before export"),
                                           ("ui-export.png", "Export", "Technical check and the \u201cdoes the story follow the plan?\u201d check")]):
        x = 0.6 + i * 4.1
        picture(s, img, x, 1.95, 3.95)
        txt(s, x, 4.35, 3.95, 0.4, head, 20, True, "green", SERIF)
        txt(s, x, 4.8, 3.95, 0.9, desc, 14, False, "muted")
    s_box = box(s, 4.4, 5.95, 4.5, 0.75, "white", "gold", lw=3)
    txt(s, 4.4, 5.95, 4.5, 0.75, "Nothing is published automatically", 18, True, "darkgold", SERIF, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

    # 6 Market
    s = d.slide("Market and users", "Who needs it, and where it fits", ("market", 0, 6),
                "Sources: Grand View Research, AI Video Generator Market 2026\u20132033; Wyzowl 2026; vendor websites (OpusClip, Descript, Synthesia, Panopto). Positioning is our judgement, not a benchmark.",
                "On screen: three target users, market growth, then the landscape map.")
    txt(s, 0.6, 1.8, 3.3, 0.3, "TARGET USERS", 12, True, "green")
    for i, (head, sub) in enumerate([("University teaching and communications teams", "Lectures to short explainers"),
                                     ("Corporate learning and development", "Training and briefings"),
                                     ("Research communicators", "Findings for the public")]):
        y = 2.15 + i * 1.5
        card(s, 0.6, y, 3.3, 1.35)
        txt(s, 0.8, y + 0.15, 2.9, 0.8, head, 16, True)
        txt(s, 0.8, y + 0.95, 2.9, 0.3, sub, 12, False, "muted")
    txt(s, 4.25, 1.8, 4.3, 0.3, "MARKET: AI VIDEO GENERATORS", 12, True, "green")
    for y, yr, share, label in [(2.2, "2025", 0.23, "USD 788.5M"), (3.05, "2033", 1.0, "USD 3,441.6M")]:
        txt(s, 4.25, y, 0.7, 0.6, yr, 16, True, anchor=MSO_ANCHOR.MIDDLE)
        box(s, 5.0, y, 3.55, 0.6, "track", None, radius=0.15)
        box(s, 5.0, y, 3.55 * share, 0.6, "green", None, radius=0.15)
        if share < 0.5:
            txt(s, 5.0 + 3.55 * share + 0.12, y, 2.5, 0.6, label, 16, True, "green", SERIF, anchor=MSO_ANCHOR.MIDDLE)
        else:
            txt(s, 5.0, y, 3.4, 0.6, label, 16, True, "white", SERIF, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)
    pill(s, 4.25, 3.95, 3.2, 0.5, "20.3% a year (2026\u20132033)", "mint", "mintline", "green", 15)
    card(s, 4.25, 4.75, 4.3, 1.75)
    txt(s, 4.45, 4.85, 1.6, 1.0, "63%", 40, True, "darkgold", SERIF)
    txt(s, 4.45, 5.75, 3.9, 0.7, "of video marketers already use AI to make or edit video (51% a year earlier)", 13, False, "muted")
    txt(s, 8.9, 1.8, 3.8, 0.3, "LANDSCAPE", 12, True, "green")
    cx, cy, cw, ch = 8.9, 2.15, 3.83, 3.75
    box(s, cx, cy, cw, ch, "white", "ink", 1.5, rounded=False)
    box(s, cx + cw / 2 - 0.01, cy, 0.02, ch, "line", rounded=False)
    box(s, cx, cy + ch / 2 - 0.01, cw, 0.02, "line", rounded=False)
    txt(s, cx + 0.08, cy + 0.05, 2.0, 0.25, "Source-grounded learning", 10, True, "muted")
    txt(s, cx + 0.08, cy + ch - 0.3, 2.0, 0.25, "Social virality", 10, True, "muted")
    txt(s, cx, cy + ch + 0.05, 1.9, 0.25, "\u2190 Fully automatic", 10, True, "muted")
    txt(s, cx + cw - 1.9, cy + ch + 0.05, 1.9, 0.25, "Human-approved \u2192", 10, True, "muted", align=PP_ALIGN.RIGHT)
    for name, fx, fy, me in [("OpusClip", .06, .78, False), ("Descript", .52, .66, False), ("Synthesia", .08, .40, False), ("Panopto", .48, .44, False), ("Brevio Studio", .52, .06, True)]:
        w, h = (1.7, 0.42) if me else (1.0, 0.32)
        pill(s, cx + 0.05 + fx * (cw - w - 0.1), cy + fy * (ch - h - 0.1), w, h, name,
             "green" if me else "white", "green" if me else "mintline", "white" if me else "ink", 13 if me else 11)
    txt(s, cx, cy + ch + 0.35, cw, 0.3, "Our positioning from vendor descriptions. Not a benchmark.", 10, False, "muted", italic=True)

    # 7 Value
    s = d.slide("Value proposition", "Faster. Consistent. Human-led.", ("value", 0, 3),
                "Sources: TechSmith, 2026 Video Statistics (87% prefer a human speaker); Brevio baseline run and 15 Oct 2026 validation (docs/four-week-plan.md, docs/project-overview.html).",
                "On screen: three pillars; each number appears with its sentence.")
    for i, (colour, head, big, text, note) in enumerate([
            ("green", "Faster highlights", "\u22485 min", "of local processing in our test run on a 15-minute lecture, before human review.", "Time saved versus manual editing is not yet measured."),
            ("gold", "Consistent quality", "14", "technical measures plus a \u201cstory follows the plan\u201d check on every export.", ""),
            ("red", "A human in control", "87%", "of viewers still prefer a human speaker, so the expert\u2019s own voice and words stay central.", "A named reviewer approves before export.")]):
        x = 0.6 + i * 4.1
        box(s, x, 1.85, 3.9, 0.09, colour, rounded=False)
        txt(s, x, 2.1, 3.9, 0.6, head, 24, True, "ink", SERIF)
        txt(s, x, 2.9, 3.9, 1.2, big, 60, True, {"green": "green", "gold": "darkgold", "red": "red"}[colour], SERIF)
        txt(s, x, 4.3, 3.7, 1.3, text, 17)
        if note:
            txt(s, x, 5.75, 3.7, 0.7, note, 13, False, "muted")

    # 8 Feasibility
    s = d.slide("Feasibility", "It already runs end to end", ("feasibility", 0, 5),
                "Internal measurements: baseline 919.6 s lecture (docs/four-week-plan.md); validation 15 Oct 2026 (docs/project-overview.html). Stage times exclude cloud storyboard, export and human review.",
                "On screen: measured stage times, three results, then the local versus cloud split.")
    txt(s, 0.6, 1.8, 12.1, 0.4, [[("919.6 s test lecture \u2192 ", {}), ("298.6 s", {"bold": True, "color": "ink"}), (" of local processing", {})]], 18, False, "muted")
    x, total = 0.6, 12.13
    for name, secs, colour in [("Transcribe", 120.5, "green"), ("Nominate", 111.5, "slate"), ("", 12.0, "gold"), ("Preview", 54.5, "red")]:
        w = total * secs / 298.6
        seg = box(s, x, 2.3, w - 0.04, 0.8, colour, None, radius=0.1)
        if name:
            txt(s, x, 2.3, w - 0.04, 0.8, name, 16, True, "white", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        txt(s, x, 3.15, w - 0.04, 0.3, f"{secs:g} s", 12, False, "muted", align=PP_ALIGN.CENTER)
        x += w
    for i, (big, small) in enumerate([("3 audiences", "Researcher, general public, kids"), ("37.6\u201345.6 s", "of finished 1080p video per run"), ("14 checks", "technical measures passed")]):
        xx = 0.6 + i * 4.1
        card(s, xx, 3.75, 3.9, 1.35)
        txt(s, xx + 0.25, 3.85, 3.5, 0.7, big, 30, True, "green", SERIF)
        txt(s, xx + 0.25, 4.6, 3.5, 0.4, small, 14, False, "muted")
    card(s, 0.6, 5.4, 7.2, 1.3, "mint", "mintline")
    txt(s, 0.85, 5.52, 6.7, 1.1, [[("On your machine", {"bold": True, "size": 17, "color": "green"})], "Transcription, voices and rendering"], 16)
    card(s, 8.05, 5.4, 4.68, 1.3, "goldtint", "goldline")
    txt(s, 8.3, 5.52, 4.2, 1.1, [[("Cloud gateway", {"bold": True, "size": 17, "color": "darkgold"})], "Short text requests only, at most 5 per run"], 16)

    # 9 Risk
    s = d.slide("Risk awareness", "What could go wrong, and what we do", ("risk", 0, 5),
                "Sources: Brevio 15 Oct 2026 validation snapshot (docs/project-overview.html); README cloud setup notes (pricing configuration is operator-accepted and unverified).",
                "Say the failures plainly; they are in the checks, not hidden.")
    rows = [("Model summaries can be wrong", "Every claim tied to a quote. Named human approval before export.", ""),
            ("Caption reading speed failed on lecture excerpts", "Measured, not hidden: reported in every check.", "On the fix list"),
            ("Some older renderer tests still fail", "Seven assertions were open at the 15 Oct snapshot.", "On the fix list"),
            ("No user trial yet", "No time-saving claim. The pilot measures it.", ""),
            ("Reuse rights and a synthetic voice", "Rights-cleared lectures only; voice disclosed.", ""),
            ("Cloud gateway access and pricing unverified", "Capped requests; reviewed pricing before the pilot.", "")]
    for i, (risk, response, status) in enumerate(rows):
        y = 1.8 + i * 0.84
        box(s, 0.6, y + 0.78, 12.13, 0.02, "line", rounded=False)
        txt(s, 0.6, y, 4.9, 0.74, risk, 17, True, anchor=MSO_ANCHOR.MIDDLE)
        txt(s, 5.6, y, 0.4, 0.74, "\u2192", 20, False, "muted", anchor=MSO_ANCHOR.MIDDLE)
        txt(s, 6.05, y, 4.9, 0.74, response, 15, False, "muted", anchor=MSO_ANCHOR.MIDDLE)
        if status:
            pill(s, 11.2, y + 0.18, 1.5, 0.38, status, "redtint", "redline", "red", 11)

    # 10 CTA
    s = d.slide("Call to action", "Help us prove it in a real pilot", ("cta", 0, 5),
                "The ask is illustrative, for a course assessment. R&D Tax Incentive: business.gov.au (current rules until 1 July 2028). Voice-over is synthetic (local Kokoro).",
                "Close on the last line and stop; do not add new claims.")
    asks = [("green", "FUNDING BODIES", "Co-fund a six-month supervised pilot", "Eligible work may qualify for the R&D Tax Incentive.", "Illustrative: AUD 150,000"),
            ("gold", "INDUSTRY PARTNERS", "Bring two or three real lectures", "Rights-cleared sources, a named reviewer and honest feedback.", ""),
            ("red", "INVESTORS", "A small seed to move from proof of concept to pilots", "Supervised, measured, no auto-publishing.", "")]
    for i, (colour, who, head, text, tag) in enumerate(asks):
        x = 0.6 + i * 4.1
        card(s, x, 1.85, 3.9, 3.3)
        box(s, x, 1.85, 3.9, 0.09, colour, rounded=False)
        txt(s, x + 0.3, 2.15, 3.3, 0.3, who, 12, True, "muted")
        txt(s, x + 0.3, 2.5, 3.35, 1.3, head, 22, True, "ink", SERIF)
        txt(s, x + 0.3, 3.85, 3.35, 0.7, text, 14)
        if tag:
            pill(s, x + 0.3, 4.55, 2.6, 0.38, tag, "goldtint", "goldline", "ink", 12)
    txt(s, 0.6, 5.5, 2.0, 0.5, "We will measure:", 18, True, anchor=MSO_ANCHOR.MIDDLE)
    for i, name in enumerate(["Time saved", "Reviewer approvals", "Audience understanding"]):
        pill(s, 2.7 + i * 2.9, 5.5, 2.7, 0.5, name, "mint", "mintline", "green", 16)
    txt(s, 0.6, 6.15, 12.1, 0.8, "Try Brevio on one lecture.", 38, True, "green", SERIF)

    # 11 References
    s = d.slide("References (appendix)", "Sources behind the numbers", None, "", "Appendix only; not part of the 4-minute run.")
    refs = ["Guo, P. J., Kim, J. & Rubin, R. (2014). How video production affects student engagement: an empirical study of MOOC videos. L@S \u201914. doi 10.1145/2556325.2566239",
            "Wyzowl (2026). State of Video Marketing 2026. 266 respondents, fieldwork late 2025. wyzowl.com/video-marketing-statistics",
            "TechSmith (2026). 2026 Video Statistics: Key Viewer Insights. techsmith.com/blog/video-statistics",
            "Grand View Research (2026). AI Video Generator Market Size, Share & Trends Analysis Report, 2026\u20132033.",
            "Vendor websites for the landscape view: opus.pro, descript.com, synthesia.io, panopto.com (positioning is our judgement).",
            "Australian Government, business.gov.au: R&D Tax Incentive (current rules until 1 July 2028).",
            "Internal: docs/four-week-plan.md (baseline run) and docs/project-overview.html (15 Oct 2026 validation snapshot).",
            "The funding ask (AUD 150,000, six-month pilot) is illustrative for a course assessment. The voice-over is synthetic."]
    txt(s, 0.6, 1.8, 12.1, 4.9, [f"\u2022  {r}" for r in refs], 15, False, "ink")

    out = PITCH / "brevio-pitch.pptx"
    d.prs.save(out)
    return out


if __name__ == "__main__":
    print(build())
